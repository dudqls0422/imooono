"""`해석결과 확인_EX.py` 테스트 (Plan 지정 12건 + 원본 무변경 확인).

파일명에 공백이 있어 import 문으로 못 불러오므로 importlib 로 로드한다.
실행: python -m unittest
"""
import contextlib
import hashlib
import importlib.util
import os
import shutil
import tempfile
import unittest
from unittest import mock

from openpyxl import load_workbook
from PIL import Image as PILImage

HERE = os.path.dirname(os.path.abspath(__file__))
MOD_PATH = os.path.join(HERE, "해석결과 확인_EX.py")
ORIGINAL_PATH = os.path.join(HERE, "해석결과 확인.py")
DESKTOP_ORIGINAL = r"C:\Users\KYB\Desktop\해석결과 캡쳐\해석결과 확인.py"


def load_module():
    spec = importlib.util.spec_from_file_location("analysis_result_ex", MOD_PATH)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


mod = load_module()


def make_jpg(path, size=(1280, 768), color=(200, 120, 60)):
    PILImage.new("RGB", size, color).save(path, "JPEG")


def _sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


class TmpMixin(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="test_arex_")
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        # 뮤테이트하는 모듈 전역 스냅샷 후 복원
        self._saved = {k: getattr(mod, k) for k in
                       ("EXPORT_DIR", "MATERIAL", "INCLUDE_TRUSS_FORCE", "headers")}
        self.addCleanup(self._restore)

    def _restore(self):
        for k, v in self._saved.items():
            setattr(mod, k, v)

    @contextlib.contextmanager
    def mock_requests(self, fail_capture_indices=()):
        """requests.post/put 를 목으로. CAPTURE 는 EXPORT_PATH 에 실제 jpg 를 써서
        '캡처가 파일을 만든 것'처럼 흉내낸다. fail_capture_indices 의 순번은 비200 처리."""
        state = {"cap": 0}

        def post(url, headers=None, json=None, **kw):
            resp = mock.Mock()
            resp.text = "{}"
            resp.json = lambda: {}
            if url.endswith("/view/CAPTURE"):
                idx = state["cap"]
                state["cap"] += 1
                if idx in fail_capture_indices:
                    resp.status_code = 500
                else:
                    resp.status_code = 200
                    make_jpg(json["Argument"]["EXPORT_PATH"])
            else:
                resp.status_code = 200
            return resp

        def put(url, headers=None, json=None, **kw):
            resp = mock.Mock()
            resp.status_code = 200
            resp.text = "{}"
            return resp

        with mock.patch.object(mod, "requests") as m:
            m.post.side_effect = post
            m.put.side_effect = put
            yield m

    def run_main_with_export_dir(self, export_dir, **mock_kw):
        mod.EXPORT_DIR = export_dir
        mod.MATERIAL = "STL"
        mod.INCLUDE_TRUSS_FORCE = True
        mod.headers = {"Content-Type": "application/json", "MAPI-Key": "dummy"}
        with self.mock_requests(**mock_kw), \
                mock.patch.object(mod, "messagebox"), \
                mock.patch.object(mod.os, "startfile", create=True):
            mod.main()


# ---------------------------------------------------------------------------
# build_result_excel (1-6)
# ---------------------------------------------------------------------------
class BuildResultExcelTests(TmpMixin):
    def _open(self, out):
        return load_workbook(out)["해석결과"]

    def test_01_images_count_matches_paths(self):
        p1 = os.path.join(self.tmp, "a.jpg")
        p2 = os.path.join(self.tmp, "b.jpg")
        make_jpg(p1)
        make_jpg(p2)
        items = [(p1, "캡션1"), (p2, "캡션2"), (p1, "캡션3")]
        out = os.path.join(self.tmp, "out1.xlsx")
        mod.build_result_excel(items, out, "STL")
        ws = self._open(out)
        self.assertEqual(len(ws._images), 3)

    def test_02_none_path_placeholder(self):
        p1 = os.path.join(self.tmp, "a.jpg")
        make_jpg(p1)
        items = [(p1, "성공"), (None, "실패한 스텝")]
        out = os.path.join(self.tmp, "out2.xlsx")
        mod.build_result_excel(items, out, "RC")
        ws = self._open(out)
        self.assertEqual(len(ws._images), 1)
        texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertTrue(any("캡처 실패" in t for t in texts))

    def test_03_page_setup_landscape_a4(self):
        p1 = os.path.join(self.tmp, "a.jpg")
        make_jpg(p1)
        out = os.path.join(self.tmp, "out3.xlsx")
        mod.build_result_excel([(p1, "c1"), (p1, "c2")], out, "STL")
        ws = self._open(out)
        self.assertEqual(int(ws.page_setup.paperSize), 9)
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(int(ws.page_setup.fitToWidth), 1)

    def test_04_caption_cell_text_exact(self):
        p1 = os.path.join(self.tmp, "a.jpg")
        make_jpg(p1)
        cap = "가새 인장력(ENV_STR)"
        out = os.path.join(self.tmp, "out4.xlsx")
        mod.build_result_excel([(p1, cap), (p1, "다음")], out, "STL")
        ws = self._open(out)
        self.assertEqual(ws["A1"].value, cap)

    def test_05_page_break_per_item(self):
        p1 = os.path.join(self.tmp, "a.jpg")
        make_jpg(p1)
        items = [(p1, f"c{i}") for i in range(4)]
        out = os.path.join(self.tmp, "out5.xlsx")
        mod.build_result_excel(items, out, "STL")
        ws = self._open(out)
        self.assertGreaterEqual(len(ws.row_breaks), len(items) - 1)

    def test_06_empty_items_raises(self):
        out = os.path.join(self.tmp, "out6.xlsx")
        with self.assertRaises(ValueError):
            mod.build_result_excel([], out, "STL")


# ---------------------------------------------------------------------------
# build_steps / main 통합 (7-12)
# ---------------------------------------------------------------------------
class StepsAndMainTests(TmpMixin):
    def test_07_build_steps_include_truss(self):
        steps = mod.build_steps(True)
        self.assertEqual(len(steps), 15)
        self.assertIn(mod.TRUSS_FORCE, steps)

    def test_08_build_steps_exclude_truss(self):
        steps = mod.build_steps(False)
        self.assertEqual(len(steps), 14)
        self.assertNotIn(mod.TRUSS_FORCE, steps)

    def test_09_main_leaves_only_xlsx_in_export_dir(self):
        d = os.path.join(self.tmp, "export")
        os.makedirs(d)
        for i in range(1, 4):
            make_jpg(os.path.join(d, f"model_{i:03d}.jpg"))
        existing = sorted(f for f in os.listdir(d) if f.startswith("model_"))

        self.run_main_with_export_dir(d)

        files = os.listdir(d)
        new_jpgs = [f for f in files if f.lower().endswith(".jpg") and f not in existing]
        self.assertEqual(new_jpgs, [])
        self.assertIn(mod.RESULT_XLSX_NAME, files)
        still = sorted(f for f in os.listdir(d) if f.startswith("model_"))
        self.assertEqual(still, existing)

    def test_10_main_removes_tempdir_on_success(self):
        d = os.path.join(self.tmp, "export10")
        os.makedirs(d)
        created = {}
        real_mkdtemp = tempfile.mkdtemp

        def rec(*a, **k):
            p = real_mkdtemp(*a, **k)
            created["p"] = p
            return p

        with mock.patch.object(mod.tempfile, "mkdtemp", side_effect=rec):
            self.run_main_with_export_dir(d)
        self.assertIn("p", created)
        self.assertFalse(os.path.exists(created["p"]))

    def test_11_main_removes_tempdir_when_excel_fails(self):
        d = os.path.join(self.tmp, "export11")
        os.makedirs(d)
        created = {}
        real_mkdtemp = tempfile.mkdtemp

        def rec(*a, **k):
            p = real_mkdtemp(*a, **k)
            created["p"] = p
            return p

        mod.EXPORT_DIR = d
        mod.MATERIAL = "STL"
        mod.INCLUDE_TRUSS_FORCE = True
        mod.headers = {"MAPI-Key": "dummy"}
        with mock.patch.object(mod.tempfile, "mkdtemp", side_effect=rec), \
                mock.patch.object(mod, "build_result_excel", side_effect=RuntimeError("boom")), \
                self.mock_requests(), \
                mock.patch.object(mod, "messagebox"), \
                mock.patch.object(mod.os, "startfile", create=True):
            with self.assertRaises(RuntimeError):
                mod.main()
        self.assertFalse(os.path.exists(created["p"]))

    def test_12_one_capture_non200_placeholder_rest_ok(self):
        d = os.path.join(self.tmp, "export12")
        os.makedirs(d)
        self.run_main_with_export_dir(d, fail_capture_indices=(2,))
        out = os.path.join(d, mod.RESULT_XLSX_NAME)
        self.assertTrue(os.path.exists(out))
        ws = load_workbook(out)["해석결과"]
        texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertTrue(any("캡처 실패" in t for t in texts))
        # 15 스텝 중 1건만 실패 → 최소 13장은 임베드
        self.assertGreaterEqual(len(ws._images), 13)


# ---------------------------------------------------------------------------
# 원본 무변경 확인
# ---------------------------------------------------------------------------
class OriginalUntouchedTests(unittest.TestCase):
    def test_repo_original_matches_desktop_original(self):
        self.assertTrue(os.path.exists(ORIGINAL_PATH))
        if os.path.exists(DESKTOP_ORIGINAL):
            self.assertEqual(_sha(ORIGINAL_PATH), _sha(DESKTOP_ORIGINAL))

    def test_original_has_no_excel_code(self):
        with open(ORIGINAL_PATH, encoding="utf-8") as f:
            src = f.read()
        self.assertNotIn("openpyxl", src)
        self.assertNotIn("build_result_excel", src)


if __name__ == "__main__":
    unittest.main()
