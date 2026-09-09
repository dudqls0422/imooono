"""Excel 입력 폼 읽기 + 검증 + 폼 생성(``--make-template``).

레이아웃은 표가 아니라 **폼(모눈 스타일)**: 시트 ``입력``, 제목 바, 섹션 5개
([1.형상]→[2.프레임 모드]→[3.재료]→[4.단면]→[5.지점]), 필드마다 라벨 박스 +
입력 박스(비어 있음) + 예시 열(회색 이탤릭).

단일 출처: ``FIELDS`` dict (key → (row, col, defined_name)). 생성기·파서·
테스트 픽스처가 전부 여기서 파생된다. 좌표를 바꾸면 이 dict 만 고치고
생성기·파서·테스트를 함께 갱신한다.

파서는 값을 먼저 **정의된 이름**으로 찾고, 없으면 ``FIELDS`` 의 고정 셀로
폴백(warning)하며, 둘 다 비어 있으면 필드명·셀 주소를 담아 오류.

재료·단면명은 **MIDAS DB(현재 모델)에 이미 정의돼 있어야 한다** — 이 도구는
만들지 않고 이름으로 기존 id 를 찾는다(``resolver``).
"""
from __future__ import annotations

import math
import warnings
from dataclasses import dataclass, field
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

# ==========================================================================
# 단일 출처 — 필드 좌표 + 정의된 이름 (col: A=1, F=6)
# ==========================================================================
SHEET_NAME = "입력"
FIELDS = {
    "span":           (5, 6, "span"),
    "eave_height":    (6, 6, "eave_height"),
    "roof_pitch":     (7, 6, "roof_pitch"),
    "bays":           (8, 6, "bays"),
    "bay_spacing":    (9, 6, "bay_spacing"),
    "frame_mode":     (12, 6, "frame_mode"),
    "material":       (15, 6, "material"),
    "sec_column":     (18, 6, "sec_column"),
    "sec_rafter":     (19, 6, "sec_rafter"),
    "sec_eave_strut": (20, 6, "sec_eave_strut"),
    "support":        (23, 6, "support"),
}

# 예시값 (입력 박스가 아니라 별도 예시 열에 표기)
EXAMPLE_VALUES = {
    "span": 28, "eave_height": 6, "roof_pitch": 0.15, "bays": 5,
    "bay_spacing": 6, "frame_mode": "3D", "material": "SS275",
    "sec_column": "H-400x200x8x13", "sec_rafter": "H-350x175x7x11",
    "sec_eave_strut": "H-200x100x5.5x8", "support": "pinned",
}

_LABEL = {
    "span": "스팬 (m)",
    "eave_height": "처마고 (m)",
    "roof_pitch": "지붕물매 (rise:run)",
    "bays": "베이 수",
    "bay_spacing": "베이 간격 (m)",
    "frame_mode": "프레임 모드",
    "material": "재료명",
    "sec_column": "기둥 단면명",
    "sec_rafter": "rafter 단면명",
    "sec_eave_strut": "이브 스트럿 단면명 (3D 전용)",
    "support": "지점 조건",
}

_COMMENT = {
    "span": "m, 기둥 중심선 간 거리",
    "eave_height": "m, 지점 레벨에서 이브(기둥·rafter 교점)까지 수직높이. 처마 끝 아님",
    "roof_pitch": "rise/run. 0.15 또는 3:12 / 3/12. 각도 아님. (0.15 = 1.5:10)",
    "bays": "정수 ≥ 1",
    "bay_spacing": "m, 프레임 간 거리",
    "frame_mode": "2D 또는 3D 만",
    "material": "MIDAS 재료 DB에 이미 존재해야 함. 정확히 일치",
    "sec_column": "MIDAS 단면 DB에 이미 존재해야 함. 대소문자·공백 정확히 일치",
    "sec_rafter": "MIDAS 단면 DB에 이미 존재해야 함. 대소문자·공백 정확히 일치",
    "sec_eave_strut": "3D 전용 — 2D에서는 무시됨",
    "support": "pinned 또는 fixed 만",
}

# (헤더 행, 제목)
_SECTIONS = [
    (4, "[1. 형상]"),
    (11, "[2. 프레임 모드]"),
    (14, "[3. 재료 — MIDAS DB에 미리 정의]"),
    (17, "[4. 단면 — MIDAS DB에 미리 정의]"),
    (22, "[5. 지점]"),
]

_GRID_COLS = 14          # A..N
_GRID_WIDTH = 3.4
_EXAMPLE_COL = 10        # J

_FIXITY = {"pinned", "fixed"}
_MODE = {"2D", "3D"}

# 지붕 경사각(도) 경고/거부 임계
_ANGLE_WARN = 30.0
_ANGLE_REJECT = 45.0


class TemplateError(ValueError):
    """템플릿 파싱/검증 실패 (거부)."""


@dataclass
class TemplateParams:
    span_m: float
    eave_height_m: float
    pitch_rise: float
    pitch_run: float
    roof_angle_deg: Optional[float]
    bay_count: int
    bay_spacing_m: float
    base_level_m: float
    column_section: str
    rafter_section: str
    eave_strut_section: Optional[str]
    material_name: str
    base_fixity: str
    frame_mode: str
    warnings: List[str] = field(default_factory=list)

    def roof_angle(self) -> float:
        if self.roof_angle_deg is not None:
            return self.roof_angle_deg
        return math.degrees(math.atan2(self.pitch_rise, self.pitch_run))


# ==========================================================================
# 변환 헬퍼
# ==========================================================================
def _to_float(v, name):
    if v is None or v == "":
        raise TemplateError(f"'{name}' 값이 비어 있음")
    try:
        return float(v)
    except (TypeError, ValueError):
        raise TemplateError(f"'{name}' 값이 숫자가 아님: {v!r}")


def _to_int(v, name):
    fv = _to_float(v, name)
    if fv != int(fv):
        raise TemplateError(f"'{name}' 는 정수여야 함: {v!r}")
    return int(fv)


def _to_str(v, name):
    s = "" if v is None else str(v).strip()
    if not s:
        raise TemplateError(f"'{name}' 값이 비어 있음 (단면/재료명은 필수)")
    return s


def _parse_pitch(v):
    """rise:run / rise/run 비율 문자열 또는 십진 slope → slope(float)."""
    if isinstance(v, str) and (":" in v or "/" in v):
        sep = ":" if ":" in v else "/"
        a, _, b = v.partition(sep)
        try:
            slope = float(a.strip()) / float(b.strip())
        except (ValueError, ZeroDivisionError):
            raise TemplateError(f"'roof_pitch' rise:run 파싱 실패: {v!r}")
    else:
        try:
            slope = float(v)
        except (TypeError, ValueError):
            raise TemplateError(f"'roof_pitch' 값이 숫자/비율이 아님: {v!r}")
    if slope <= 0:
        raise TemplateError(f"'roof_pitch' 는 0 보다 커야 함 (slope={slope})")
    return slope


# ==========================================================================
# 읽기 — 정의된 이름 우선, 고정 셀 폴백
# ==========================================================================
def _from_defined_name(wb, dn):
    d = wb.defined_names.get(dn)
    if d is None:
        return None, False
    dests = list(d.destinations)
    if not dests:
        return None, False
    sheet, coord = dests[0]
    try:
        return wb[sheet][coord].value, True
    except KeyError:
        return None, False


def _raw(wb, ws, key, *, required=True):
    row, col, dn = FIELDS[key]
    a1 = f"{get_column_letter(col)}{row}"
    val, ok = _from_defined_name(wb, dn)
    if not ok:
        val = ws.cell(row=row, column=col).value
        warnings.warn(
            f"'{key}': 정의된 이름 '{dn}' 없음 — 대체 셀 {SHEET_NAME}!{a1} 사용",
            stacklevel=3)
    if val is None or (isinstance(val, str) and val.strip() == ""):
        if not required:
            return None
        raise TemplateError(
            f"'{key}' 입력값 없음 (정의된 이름 '{dn}' 없음/빈값, "
            f"대체 셀 {a1} 비어 있음)")
    return val


def read_template(path: str) -> TemplateParams:
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    def f(key, **kw):
        return _raw(wb, ws, key, **kw)

    frame_mode = str(f("frame_mode")).strip().upper()
    if frame_mode not in _MODE:
        raise TemplateError(f"'frame_mode' 값은 2D|3D 여야 함: {frame_mode!r}")
    support = str(f("support")).strip().lower()
    if support not in _FIXITY:
        raise TemplateError(f"'support' 값은 pinned|fixed 여야 함: {support!r}")

    slope = _parse_pitch(f("roof_pitch"))

    if frame_mode == "3D":
        eave_strut = _to_str(f("sec_eave_strut"), "sec_eave_strut")
    else:
        raw_es = f("sec_eave_strut", required=False)
        eave_strut = (str(raw_es).strip()
                      if raw_es not in (None, "") else None)

    params = TemplateParams(
        span_m=_to_float(f("span"), "span"),
        eave_height_m=_to_float(f("eave_height"), "eave_height"),
        pitch_rise=slope, pitch_run=1.0, roof_angle_deg=None,
        bay_count=_to_int(f("bays"), "bays"),
        bay_spacing_m=_to_float(f("bay_spacing"), "bay_spacing"),
        base_level_m=0.0,
        column_section=_to_str(f("sec_column"), "sec_column"),
        rafter_section=_to_str(f("sec_rafter"), "sec_rafter"),
        eave_strut_section=eave_strut,
        material_name=_to_str(f("material"), "material"),
        base_fixity=support,
        frame_mode=frame_mode,
    )
    validate(params)
    return params


# ==========================================================================
# 검증 (기능·규칙 무변경)
# ==========================================================================
def validate(p: TemplateParams) -> TemplateParams:
    if p.span_m <= 0:
        raise TemplateError(f"span_m 은 0 보다 커야 함: {p.span_m}")
    if p.eave_height_m <= 0:
        raise TemplateError(f"eave_height_m 은 0 보다 커야 함: {p.eave_height_m}")
    if p.bay_count < 1:
        raise TemplateError(f"bay_count 는 1 이상이어야 함: {p.bay_count}")
    if p.bay_spacing_m <= 0:
        raise TemplateError(f"bay_spacing_m 은 0 보다 커야 함: {p.bay_spacing_m}")
    if p.base_fixity not in _FIXITY:
        raise TemplateError(f"base_fixity 는 pinned|fixed: {p.base_fixity!r}")
    if p.frame_mode not in _MODE:
        raise TemplateError(f"frame_mode 는 2D|3D: {p.frame_mode!r}")

    if p.roof_angle_deg is None:
        if p.pitch_rise <= 0 or p.pitch_run <= 0:
            raise TemplateError(
                f"pitch_rise/pitch_run 은 0 보다 커야 함: {p.pitch_rise}/{p.pitch_run}")
    else:
        if p.roof_angle_deg <= 0:
            raise TemplateError(f"roof_angle_deg 는 0 보다 커야 함: {p.roof_angle_deg}")

    angle = p.roof_angle()
    if angle >= _ANGLE_REJECT:
        raise TemplateError(
            f"지붕 경사각 {angle:.1f}° — {_ANGLE_REJECT:.0f}° 이상은 포탈 프레임 범위 밖(거부)")
    if angle >= _ANGLE_WARN:
        p.warnings.append(
            f"지붕 경사각 {angle:.1f}° — {_ANGLE_WARN:.0f}~{_ANGLE_REJECT:.0f}° 는 이례적(경고)")

    slope = (math.tan(math.radians(p.roof_angle_deg))
             if p.roof_angle_deg is not None else p.pitch_rise / p.pitch_run)
    ridge = p.base_level_m + p.eave_height_m + (p.span_m / 2.0) * slope
    eave = p.base_level_m + p.eave_height_m
    if ridge <= eave + 1e-9:
        raise TemplateError(
            f"용마루({ridge:.4f} m) ≤ 처마({eave:.4f} m) — 슬로프가 0 이하")
    return p


# ==========================================================================
# 폼 생성 (--make-template)
# ==========================================================================
_TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
_TITLE_FONT = Font(bold=True, color="FFFFFF", size=12)
_SEC_FILL = PatternFill("solid", fgColor="D9E2F3")
_SEC_FONT = Font(bold=True, size=10)
_LABEL_FILL = PatternFill("solid", fgColor="F2F2F2")
_LABEL_FONT = Font(bold=True, size=10)
_INPUT_FILL = PatternFill("solid", fgColor="FFFFFF")
_GREY_FILL = PatternFill("solid", fgColor="E8E8E8")
_EX_FONT = Font(italic=True, color="808080", size=9)
_THIN = Side(style="thin", color="AAAAAA")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MID = Alignment(horizontal="left", vertical="center")


def _sheet_ref(row, col):
    return f"'{SHEET_NAME}'!${get_column_letter(col)}${row}"


def write_template(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    for c in range(1, _GRID_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = _GRID_WIDTH

    # 제목 바 A1:N2
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=_GRID_COLS)
    t = ws.cell(row=1, column=1,
                value="포탈 프레임 모델 생성 입력 시트   (단위: m, kN)")
    t.font = _TITLE_FONT
    t.alignment = Alignment(horizontal="center", vertical="center")
    for r in (1, 2):
        for c in range(1, _GRID_COLS + 1):
            ws.cell(row=r, column=c).fill = _TITLE_FILL

    # 예시 열 헤더
    eh = ws.cell(row=4, column=_EXAMPLE_COL, value="예시")
    eh.font = _EX_FONT

    # 섹션 헤더 바 (A:N)
    for row, title in _SECTIONS:
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=_GRID_COLS)
        s = ws.cell(row=row, column=1, value=title)
        s.font = _SEC_FONT
        for c in range(1, _GRID_COLS + 1):
            ws.cell(row=row, column=c).fill = _SEC_FILL

    # 필드 행: 라벨 박스(A:D) + 입력 박스(F, 단일) + 예시(J) + 정의된 이름
    for key, (row, col, dn) in FIELDS.items():
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        lab = ws.cell(row=row, column=1, value=_LABEL[key])
        lab.font = _LABEL_FONT
        lab.fill = _LABEL_FILL
        lab.alignment = _MID
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = _BOX

        inp = ws.cell(row=row, column=col)
        inp.border = _BOX
        inp.fill = _GREY_FILL if key == "sec_eave_strut" else _INPUT_FILL
        inp.alignment = _MID
        cm = Comment(_COMMENT[key], "portal_frame_gen")
        cm.width, cm.height = 220, 100
        inp.comment = cm

        ex = ws.cell(row=row, column=_EXAMPLE_COL, value=EXAMPLE_VALUES[key])
        ex.font = _EX_FONT

        wb.defined_names[dn] = DefinedName(dn, attr_text=_sheet_ref(row, col))

    # 데이터 유효성 (드롭다운 — 자문용, 파서가 별도 재검증)
    dv_mode = DataValidation(type="list", formula1='"2D,3D"', allow_blank=True)
    dv_sup = DataValidation(type="list", formula1='"pinned,fixed"',
                            allow_blank=True)
    ws.add_data_validation(dv_mode)
    ws.add_data_validation(dv_sup)
    dv_mode.add(ws.cell(row=FIELDS["frame_mode"][0],
                        column=FIELDS["frame_mode"][1]))
    dv_sup.add(ws.cell(row=FIELDS["support"][0], column=FIELDS["support"][1]))

    # 페이지 설정
    last_row = FIELDS["support"][0] + 1
    ws.print_area = f"A1:{get_column_letter(_GRID_COLS)}{last_row}"
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    wb.save(path)
