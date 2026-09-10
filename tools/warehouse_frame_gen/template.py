"""창고 프레임 입력 폼(Excel) 생성 + 파싱.

portal_frame_gen 의 폼 스타일을 답습: 모눈 그리드, 병합 라벨/입력 박스,
정의된 이름, 예시 회색 열, 드롭다운, 셀 코멘트, 인쇄영역.

단일 출처 ``FIELDS`` (key → (row, col, defined_name)) 19개 + ``EXAMPLE_VALUES``.
파서는 정의된 이름 → 고정 셀 폴백(warning) → 둘 다 비면 필드명·셀을 담은
``TemplateError``. 반환은 ``build_model`` 이 그대로 쓰는 dict.
"""
from __future__ import annotations

import warnings

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from .geometry import parse_span_expression

SHEET_NAME = "입력"

# key → (row, col, defined_name).  col: A=1, F=6
FIELDS = {
    "x_spans":          (5, 6, "x_spans"),
    "ny":               (6, 6, "ny"),
    "dy":               (7, 6, "dy"),
    "nz":               (8, 6, "nz"),
    "dz":               (9, 6, "dz"),
    "ox":               (12, 6, "ox"),
    "oy":               (13, 6, "oy"),
    "oz":               (14, 6, "oz"),
    "matl":             (17, 6, "matl"),
    "sect":             (18, 6, "sect"),
    "node_start":       (19, 6, "node_start"),
    "elem_start":       (20, 6, "elem_start"),
    "skip_bottom_beam": (23, 6, "skip_bottom_beam"),
    "skip_top_beam":    (24, 6, "skip_top_beam"),
    "use_ridge":        (27, 6, "use_ridge"),
    "rise":             (28, 6, "rise"),
    "connect_rafters":  (29, 6, "connect_rafters"),
    "rafter_subdiv":    (30, 6, "rafter_subdiv"),
    "num_purlins":      (31, 6, "num_purlins"),
}

EXAMPLE_VALUES = {
    "x_spans": "2, 4@4, 2", "ny": 4, "dy": 6.0, "nz": 2, "dz": 4.0,
    "ox": 0, "oy": 0, "oz": 0,
    "matl": 1, "sect": 1, "node_start": 1, "elem_start": 1,
    "skip_bottom_beam": "TRUE", "skip_top_beam": "TRUE",
    "use_ridge": "TRUE", "rise": 1.5, "connect_rafters": "TRUE",
    "rafter_subdiv": 1, "num_purlins": 0,
}

_LABEL = {
    "x_spans": "X방향 기둥 간격 리스트",
    "ny": "Y방향 절점 개수 (ny)",
    "dy": "Y방향 간격 (dy, m)",
    "nz": "Z방향 절점 개수 (nz)",
    "dz": "Z방향 간격 (dz, m)",
    "ox": "원점 X (m)", "oy": "원점 Y (m)", "oz": "원점 Z (m)",
    "matl": "재료번호 (MATL)",
    "sect": "단면번호 (SECT)",
    "node_start": "절점 시작번호",
    "elem_start": "요소 시작번호",
    "skip_bottom_beam": "최하부 층 보 생략",
    "skip_top_beam": "지붕층 Y방향 보 생략",
    "use_ridge": "용마루 지붕 추가",
    "rise": "용마루 상승높이 (rise, m)",
    "connect_rafters": "서까래 생성",
    "rafter_subdiv": "기둥 사이 서까래 분할수",
    "num_purlins": "중도리 개수",
}

_COMMENT = {
    "x_spans": "콤마 구분, '간격@개수' 반복 지원. 예: 2, 4@4, 2 → [2,4,4,4,4,2]",
    "ny": "Y방향 절점 개수. 정수 ≥ 1",
    "dy": "Y방향 등간격 (m). > 0",
    "nz": "Z방향 절점 개수. 정수 ≥ 1. 1이면 기둥 생성 안 됨",
    "dz": "Z방향 등간격 (m). > 0",
    "ox": "원점 X 좌표 (m)", "oy": "원점 Y 좌표 (m)", "oz": "원점 Z 좌표 (m)",
    "matl": "재료번호. MIDAS 현재 모델에 이미 존재해야 함(정수)",
    "sect": "단면번호. MIDAS 현재 모델에 이미 존재해야 함(정수)",
    "node_start": "절점 시작번호. 정수 ≥ 1",
    "elem_start": "요소 시작번호. 정수 ≥ 1",
    "skip_bottom_beam": "최하부 층(k=0)에 X/Y 보 생성 안 함. 빈칸=TRUE",
    "skip_top_beam": "지붕층에 Y방향 보 생성 안 함(X방향은 용마루 옵션이 처리). 빈칸=TRUE",
    "use_ridge": "용마루(X방향 박공지붕) 추가. 빈칸=TRUE",
    "rise": "용마루 상승높이 (m). use_ridge=TRUE 면 > 0",
    "connect_rafters": "서까래(처마↔용마루) 생성. FALSE면 용마루보만. 빈칸=TRUE",
    "rafter_subdiv": "기둥 사이 서까래 분할수. 정수 ≥ 1 (1=추가 없음)",
    "num_purlins": "중도리 단수. 정수 ≥ 0 (0=없음)",
}

_SECTIONS = [
    (4, "[1. 그리드]"),
    (11, "[2. 원점]"),
    (16, "[3. 재료·단면·시작번호]"),
    (22, "[4. 보 옵션]"),
    (26, "[5. 용마루]"),
]

_BOOL_FIELDS = ("skip_bottom_beam", "skip_top_beam", "use_ridge", "connect_rafters")
_BOOL_DEFAULT = {"skip_bottom_beam": True, "skip_top_beam": True,
                 "use_ridge": True, "connect_rafters": True}
_FLOAT_FIELDS = ("dy", "dz", "ox", "oy", "oz", "rise")
_INT_FIELDS = ("ny", "nz", "matl", "sect", "node_start", "elem_start",
               "rafter_subdiv", "num_purlins")

_GRID_COLS = 14
_GRID_WIDTH = 3.4
_EXAMPLE_COL = 10

_BOOL_TRUE = {"true", "1", "y", "yes", "예", "o", "on"}
_BOOL_FALSE = {"false", "0", "n", "no", "아니오", "x", "off"}


class TemplateError(ValueError):
    """템플릿 파싱/검증 실패 (거부)."""


# ==========================================================================
# 파싱
# ==========================================================================
def _from_defined_name(wb, dn):
    d = wb.defined_names.get(dn)
    if d is None:
        return None, False
    dests = list(d.destinations)
    if not dests:
        return None, False
    sheet, coord = dests[0]
    try:
        return wb[sheet][coord].value, True
    except KeyError:
        return None, False


def _raw(wb, ws, key, *, required=True):
    row, col, dn = FIELDS[key]
    a1 = f"{get_column_letter(col)}{row}"
    val, ok = _from_defined_name(wb, dn)
    if not ok:
        val = ws.cell(row=row, column=col).value
        warnings.warn(
            f"'{key}': 정의된 이름 '{dn}' 없음 — 대체 셀 {SHEET_NAME}!{a1} 사용",
            stacklevel=3)
    if val is None or (isinstance(val, str) and val.strip() == ""):
        if not required:
            return None
        raise TemplateError(
            f"'{key}' 입력값 없음 (정의된 이름 '{dn}' 없음/빈값, "
            f"대체 셀 {a1} 비어 있음)")
    return val


def _to_float(v, name):
    try:
        return float(v)
    except (TypeError, ValueError):
        raise TemplateError(f"'{name}' 값이 숫자가 아님: {v!r}")


def _to_int(v, name):
    f = _to_float(v, name)
    if f != int(f):
        raise TemplateError(f"'{name}' 는 정수여야 함: {v!r}")
    return int(f)


def _parse_bool(v, field, default):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return default
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s in _BOOL_TRUE:
        return True
    if s in _BOOL_FALSE:
        return False
    raise TemplateError(
        f"'{field}' 불리언 값을 해석할 수 없음: {v!r} (TRUE/FALSE)")


def read_template(path: str) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

    out = {}
    for key in FIELDS:
        if key in _BOOL_FIELDS:
            raw = _raw(wb, ws, key, required=False)
            out[key] = _parse_bool(raw, key, _BOOL_DEFAULT[key])
        elif key == "x_spans":
            raw = _raw(wb, ws, key)
            try:
                out[key] = parse_span_expression(str(raw))
            except ValueError as e:
                raise TemplateError(f"'x_spans' 파싱 오류: {e}")
        elif key in _FLOAT_FIELDS:
            out[key] = _to_float(_raw(wb, ws, key), key)
        elif key in _INT_FIELDS:
            out[key] = _to_int(_raw(wb, ws, key), key)
        else:  # pragma: no cover
            out[key] = _raw(wb, ws, key)
    return out


# ==========================================================================
# 폼 생성
# ==========================================================================
_TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
_TITLE_FONT = Font(bold=True, color="FFFFFF", size=12)
_SEC_FILL = PatternFill("solid", fgColor="D9E2F3")
_SEC_FONT = Font(bold=True, size=10)
_LABEL_FILL = PatternFill("solid", fgColor="F2F2F2")
_LABEL_FONT = Font(bold=True, size=10)
_INPUT_FILL = PatternFill("solid", fgColor="FFFFFF")
_BOOL_FILL = PatternFill("solid", fgColor="FFF2CC")
_EX_FONT = Font(italic=True, color="808080", size=9)
_THIN = Side(style="thin", color="AAAAAA")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MID = Alignment(horizontal="left", vertical="center")


def _sheet_ref(row, col):
    return f"'{SHEET_NAME}'!${get_column_letter(col)}${row}"


def write_template(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    for c in range(1, _GRID_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = _GRID_WIDTH

    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=_GRID_COLS)
    t = ws.cell(row=1, column=1,
                value="창고 기본 프레임 생성 입력 시트   (단위: m, kN)")
    t.font = _TITLE_FONT
    t.alignment = Alignment(horizontal="center", vertical="center")
    for r in (1, 2):
        for c in range(1, _GRID_COLS + 1):
            ws.cell(row=r, column=c).fill = _TITLE_FILL

    ws.cell(row=4, column=_EXAMPLE_COL, value="예시").font = _EX_FONT

    for row, title in _SECTIONS:
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=_GRID_COLS)
        ws.cell(row=row, column=1, value=title).font = _SEC_FONT
        for c in range(1, _GRID_COLS + 1):
            ws.cell(row=row, column=c).fill = _SEC_FILL

    for key, (row, col, dn) in FIELDS.items():
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        lab = ws.cell(row=row, column=1, value=_LABEL[key])
        lab.font = _LABEL_FONT
        lab.fill = _LABEL_FILL
        lab.alignment = _MID
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = _BOX

        inp = ws.cell(row=row, column=col)
        inp.border = _BOX
        inp.fill = _BOOL_FILL if key in _BOOL_FIELDS else _INPUT_FILL
        inp.alignment = _MID
        cm = Comment(_COMMENT[key], "warehouse_frame_gen")
        cm.width, cm.height = 240, 96
        inp.comment = cm

        ws.cell(row=row, column=_EXAMPLE_COL,
                value=EXAMPLE_VALUES[key]).font = _EX_FONT

        wb.defined_names[dn] = DefinedName(dn, attr_text=_sheet_ref(row, col))

    dv_bool = DataValidation(type="list", formula1='"TRUE,FALSE"',
                             allow_blank=True)
    ws.add_data_validation(dv_bool)
    for key in _BOOL_FIELDS:
        r, c, _ = FIELDS[key]
        dv_bool.add(ws.cell(row=r, column=c))

    last_row = max(r for r, _c, _d in FIELDS.values()) + 1
    ws.print_area = f"A1:{get_column_letter(_GRID_COLS)}{last_row}"
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    wb.save(path)
