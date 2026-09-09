"""openpyxl 3시트 리포트: 요약 / 모델점검 / 하중조건 (+ 필요 시 오버플로)."""
from __future__ import annotations

from datetime import datetime
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .checks import Finding, SEVERITY_ERROR, SEVERITY_WARN, SEVERITY_INFO
from .loadcond import SEIS_COLS, WIND_COLS

MAX_IDS_INLINE = 50

_SEV_FILL = {
    SEVERITY_ERROR: PatternFill("solid", fgColor="F8CBAD"),
    SEVERITY_WARN: PatternFill("solid", fgColor="FFE699"),
    SEVERITY_INFO: PatternFill("solid", fgColor="DDEBF7"),
}
_HDR_FONT = Font(bold=True)
_HDR_FILL = PatternFill("solid", fgColor="D9D9D9")


def _ids_cell(ids: List) -> str:
    if not ids:
        return ""
    ids = list(ids)
    if len(ids) <= MAX_IDS_INLINE:
        return ", ".join(str(i) for i in ids)
    nums = [int(i) for i in ids if str(i).lstrip("-").isdigit()]
    span = f", ID range {min(nums)}~{max(nums)}" if len(nums) == len(ids) else ""
    return f"건수 {len(ids)}{span} (전체는 '오버플로' 시트)"


def _style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _HDR_FONT
        cell.fill = _HDR_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{max(2, ws.max_row)}"
    ws.page_setup.paperSize = 9


def _autosize(ws, ncol, cap=60):
    for c in range(1, ncol + 1):
        w = 10
        for r in range(1, min(ws.max_row, 400) + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                w = max(w, min(cap, len(str(v)) + 2))
        ws.column_dimensions[get_column_letter(c)].width = w


def _load_status(rows):
    if not rows:
        return "해당 없음(케이스 없음)"
    srcs = {str(r.get("값 출처", "")).split(" ")[0] for r in rows}
    if srcs == {"API조회"}:
        return "성공(API조회)"
    if any("간접" in s for s in srcs):
        return "간접/직접 조회 실패"
    return " / ".join(sorted(srcs))


def _write_summary(ws, findings, ctx, meta, wind_rows, seis_rows):
    active = [f for f in findings if not f.whitelisted]
    n_err = sum(1 for f in active if f.severity == SEVERITY_ERROR)
    n_warn = sum(1 for f in active if f.severity == SEVERITY_WARN)
    n_info = sum(1 for f in active if f.severity == SEVERITY_INFO)
    n_wl = sum(1 for f in findings if f.whitelisted)
    n_checks = meta.get("n_checks", 0)
    fail_ids = {f.check_id for f in active
                if f.severity in (SEVERITY_ERROR, SEVERITY_WARN)}
    n_pass = max(0, n_checks - len(fail_ids))

    plate_note = ("판요소 있음" if ctx.has_plates()
                  else "판요소 없음 — 메시 점검(M1~M3) 생략")
    cs_note = ("CS 데이터 있음 — stage-inactive 요소 제외 적용"
               if ctx.has_cs() else "CS 데이터 없음 — 일반 모델로 점검")

    rows = [
        ("스캔 일시", meta.get("timestamp", "")),
        ("base_url", meta.get("base_url", "")),
        ("모델 원단위", f"FORCE={ctx.unit_raw.get('FORCE')} DIST={ctx.unit_raw.get('DIST')}"),
        ("정규화 단위", "kN, m, 전역 XYZ, 인장 +"),
        ("", ""),
        ("심각도 — 오류", n_err),
        ("심각도 — 경고", n_warn),
        ("심각도 — 정보", n_info),
        ("화이트리스트 처리 건수", n_wl),
        ("총 점검 항목 수", n_checks),
        ("통과(오류·경고 없음) 항목 수", n_pass),
        ("", ""),
        ("절점 수", len(ctx.nodes)),
        ("요소 수", len(ctx.elems)),
        ("판요소 수", sum(1 for e in ctx.elems.values()
                       if str(e.get("TYPE", "")).upper()
                       in ("PLATE", "WALL", "MEMBRANE", "PLANESTRESS", "PLANESTRAIN"))),
        ("정적 하중케이스 수", len(ctx.stld)),
        ("하중조합 수", sum(len(v) for v in ctx.lcom.values())),
        ("층 정의 수", len(ctx.stor)),
        ("", ""),
        ("판요소", plate_note),
        ("시공단계(CS)", cs_note),
        ("풍하중 조회 상태", _load_status(wind_rows)),
        ("지진하중 조회 상태", _load_status(seis_rows)),
        ("조회 실패 리소스", ", ".join(f"db/{k}({s})" for k, s in ctx.unavailable) or "없음"),
    ]
    ws.cell(row=1, column=1, value="항목").font = _HDR_FONT
    ws.cell(row=1, column=2, value="값").font = _HDR_FONT
    for i, (k, v) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)

    r = len(rows) + 4
    ws.cell(row=r, column=1, value="상위 위험 (오류→경고 순)").font = _HDR_FONT
    top = sorted(active, key=lambda f: f.sort_key())[:5]
    for j, f in enumerate(top, start=1):
        ws.cell(row=r + j, column=1, value=f"{j}. [{f.check_id}] {f.check_name}")
        ws.cell(row=r + j, column=2, value=f"{f.severity} — {f.description}")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    ws.page_setup.paperSize = 9
    ws.freeze_panes = "A2"


def _write_findings(ws, findings, wb):
    cols = ["점검항목ID", "점검항목명", "심각도", "대상유형", "대상ID",
            "설명", "권장조치", "화이트리스트여부"]
    for c, name in enumerate(cols, start=1):
        ws.cell(row=1, column=c, value=name)
    overflow = []
    for i, f in enumerate(findings, start=2):
        ids = list(f.target_ids)
        ws.cell(row=i, column=1, value=f.check_id)
        ws.cell(row=i, column=2, value=f.check_name)
        sc = ws.cell(row=i, column=3, value=f.severity)
        if f.severity in _SEV_FILL:
            sc.fill = _SEV_FILL[f.severity]
        ws.cell(row=i, column=4, value=f.target_type)
        ws.cell(row=i, column=5, value=_ids_cell(ids))
        ws.cell(row=i, column=6, value=f.description)
        ws.cell(row=i, column=7, value=f.recommendation)
        ws.cell(row=i, column=8, value="Y" if f.whitelisted else "")
        if len(ids) > MAX_IDS_INLINE:
            overflow.append((f.check_id, ids))
    _style_header(ws, len(cols))
    _autosize(ws, len(cols))

    if overflow:
        ov = wb.create_sheet("오버플로")
        ov.cell(row=1, column=1, value="점검항목ID")
        ov.cell(row=1, column=2, value="전체 대상 ID")
        for j, (cid, ids) in enumerate(overflow, start=2):
            ov.cell(row=j, column=1, value=cid)
            ov.cell(row=j, column=2,
                    value="[" + cid + "] " + ", ".join(str(x) for x in ids))
        _style_header(ov, 2)
        ov.column_dimensions["B"].width = 120


def _write_loadcond(ws, wind_rows, seis_rows):
    r = 1
    ws.cell(row=r, column=1, value="풍하중 적용조건").font = _HDR_FONT
    r += 1
    for c, name in enumerate(WIND_COLS, start=1):
        cell = ws.cell(row=r, column=c, value=name)
        cell.font = _HDR_FONT
        cell.fill = _HDR_FILL
    for row in wind_rows:
        r += 1
        for c, name in enumerate(WIND_COLS, start=1):
            ws.cell(row=r, column=c, value=row.get(name, ""))
    if not wind_rows:
        r += 1
        ws.cell(row=r, column=1, value="(풍 하중케이스 없음)")

    r += 3
    ws.cell(row=r, column=1, value="지진하중 적용조건").font = _HDR_FONT
    r += 1
    for c, name in enumerate(SEIS_COLS, start=1):
        cell = ws.cell(row=r, column=c, value=name)
        cell.font = _HDR_FONT
        cell.fill = _HDR_FILL
    for row in seis_rows:
        r += 1
        for c, name in enumerate(SEIS_COLS, start=1):
            ws.cell(row=r, column=c, value=row.get(name, ""))
    if not seis_rows:
        r += 1
        ws.cell(row=r, column=1, value="(지진 하중케이스 없음)")

    ws.page_setup.paperSize = 9
    ws.freeze_panes = "A2"
    _autosize(ws, max(len(WIND_COLS), len(SEIS_COLS)), cap=45)


def write_report(path, findings, ctx, wind_rows, seis_rows, meta):
    meta = dict(meta or {})
    meta.setdefault("timestamp", datetime.now().isoformat(timespec="seconds"))
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "요약"
    ws_chk = wb.create_sheet("모델점검")
    ws_load = wb.create_sheet("하중조건")

    _write_findings(ws_chk, findings, wb)
    _write_loadcond(ws_load, wind_rows, seis_rows)
    _write_summary(ws_sum, findings, ctx, meta, wind_rows, seis_rows)
    wb.save(path)
