"""입력 폼 생성 + 정의된 이름/폴백 파서 + 불리언 파싱."""
from __future__ import annotations

import warnings

import pytest
from openpyxl import load_workbook

from tools.warehouse_frame_gen.template import (SHEET_NAME, TemplateError,
                                                read_template, write_template)

from .conftest import make_form


def test_roundtrip_example_values(tmp_path):
    p = read_template(make_form(tmp_path))
    assert p["x_spans"] == [2.0, 4.0, 4.0, 4.0, 4.0, 2.0]
    assert p["ny"] == 4 and isinstance(p["ny"], int)
    assert p["dy"] == 6.0 and p["nz"] == 2 and p["dz"] == 4.0
    assert p["ox"] == 0.0 and p["oy"] == 0.0 and p["oz"] == 0.0
    assert p["matl"] == 1 and p["sect"] == 1
    assert p["node_start"] == 1 and p["elem_start"] == 1
    assert p["skip_bottom_beam"] is True and p["skip_top_beam"] is True
    assert p["use_ridge"] is True and p["connect_rafters"] is True
    assert p["rise"] == 1.5 and p["rafter_subdiv"] == 1 and p["num_purlins"] == 0


def test_defined_name_missing_falls_back_with_warning(tmp_path):
    path = make_form(tmp_path, drop_names=["ny"])
    with pytest.warns(UserWarning, match="ny"):
        p = read_template(path)
    assert p["ny"] == 4


def test_missing_value_and_name_raises(tmp_path):
    path = make_form(tmp_path, drop_names=["ny"], ny=None)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with pytest.raises(TemplateError) as ei:
            read_template(path)
    assert "ny" in str(ei.value) and "F6" in str(ei.value)


def test_empty_x_spans_raises(tmp_path):
    path = make_form(tmp_path, x_spans="   ")
    with pytest.raises(TemplateError):
        read_template(path)


@pytest.mark.parametrize("raw,expect", [
    ("TRUE", True), ("true", True), ("True", True), ("1", True), ("y", True),
    ("Y", True), ("예", True), ("o", True),
    ("FALSE", False), ("false", False), ("0", False), ("n", False),
    ("아니오", False), ("x", False), ("X", False),
])
def test_bool_parsing_matrix(tmp_path, raw, expect):
    p = read_template(make_form(tmp_path, use_ridge=raw))
    assert p["use_ridge"] is expect


def test_bool_empty_uses_documented_default(tmp_path):
    p = read_template(make_form(tmp_path, skip_bottom_beam=None,
                                skip_top_beam=None, use_ridge=None,
                                connect_rafters=None))
    assert p["skip_bottom_beam"] is True and p["skip_top_beam"] is True
    assert p["use_ridge"] is True and p["connect_rafters"] is True


def test_bool_bad_value_raises(tmp_path):
    with pytest.raises(TemplateError):
        read_template(make_form(tmp_path, use_ridge="maybe"))


def test_form_has_bool_dropdown(tmp_path):
    path = tmp_path / "t.xlsx"
    write_template(str(path))
    wb = load_workbook(str(path))
    ws = wb[SHEET_NAME]
    formulas = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert '"TRUE,FALSE"' in formulas


def test_regenerate_no_duplicate_defined_names(tmp_path):
    path = tmp_path / "t.xlsx"
    write_template(str(path))
    write_template(str(path))
    wb = load_workbook(str(path))
    names = list(wb.defined_names)
    assert len(names) == len(set(names)) == 19
