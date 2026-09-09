"""창고 프레임 생성기 테스트 공용 픽스처 (라이브 MIDAS 불필요)."""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from tools.warehouse_frame_gen.geometry import parse_span_expression
from tools.warehouse_frame_gen.template import (EXAMPLE_VALUES, FIELDS,
                                                write_template)

_SKIP = object()


def make_params(**over) -> dict:
    base = dict(
        x_spans=parse_span_expression("2, 4@4, 2"),
        ny=4, dy=6.0, nz=2, dz=4.0, ox=0.0, oy=0.0, oz=0.0,
        matl=1, sect=1, node_start=1, elem_start=1,
        skip_bottom_beam=True, skip_top_beam=True,
        use_ridge=True, rise=1.5, connect_rafters=True,
        rafter_subdiv=1, num_purlins=0,
    )
    base.update(over)
    return base


def make_form(tmp_path, name="in.xlsx", drop_names=(), **overrides) -> str:
    """폼 생성 후 입력 셀(F 열)에 예시값(+오버라이드) 기입. None → 셀 비움."""
    path = tmp_path / name
    write_template(str(path))
    wb = load_workbook(str(path))
    ws = wb[wb.sheetnames[0]]
    vals = dict(EXAMPLE_VALUES)
    vals.update(overrides)
    for key, (row, col, _dn) in FIELDS.items():
        v = vals.get(key, _SKIP)
        if v is _SKIP:
            continue
        ws.cell(row=row, column=col, value=v)
    for dn in drop_names:
        if dn in wb.defined_names:
            del wb.defined_names[dn]
    wb.save(str(path))
    return str(path)


class FakeResp:
    def __init__(self, status=200, payload=None, text=""):
        self.status_code = status
        self._p = payload
        self.text = text or (json.dumps(payload) if payload is not None else "")

    def json(self):
        if self._p is None:
            raise ValueError("no json")
        return self._p


class FakeClient:
    def __init__(self, get_map=None, put_map=None):
        self.get_map = get_map or {}
        self.put_map = put_map or {}
        self.puts = []
        self.request_log = []

    def get(self, path):
        self.request_log.append(("GET", path, 200))
        return self.get_map.get(path, {"message": ""})

    def put(self, path, body):
        self.puts.append((path, body))
        self.request_log.append(("PUT", path, 200))
        key = path.rsplit("/", 1)[-1]
        return self.put_map.get(path, FakeResp(200, {key: {}}))

    def verbs_used(self):
        return sorted({v for v, _p, _s in self.request_log})


@pytest.fixture
def full_get_map():
    """matl=1, sect=1 이 정의된 빈 모델."""
    return {
        "db/MATL": {"MATL": {"1": {"NAME": "M1"}}},
        "db/SECT": {"SECT": {"1": {"SECT_NAME": "S1"}}},
        "db/NODE": {"message": ""},
        "db/ELEM": {"message": ""},
    }
