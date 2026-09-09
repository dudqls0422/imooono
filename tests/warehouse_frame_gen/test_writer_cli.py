"""writer 가드/존재확인 + CLI 종료코드 + dry-run 네트워크 0."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from tools.warehouse_frame_gen import cli
from tools.warehouse_frame_gen.geometry import build_model
from tools.warehouse_frame_gen.template import SHEET_NAME
from tools.warehouse_frame_gen.writer import (MatlSectMissing, ModelNotEmpty,
                                              check_matl_sect,
                                              empty_model_guard, put_unit,
                                              write_model)

from .conftest import FakeClient, FakeResp, make_form, make_params


def _quiet(*_a, **_k):
    pass


# --- writer: 빈 모델 가드 ------------------------------------------------
def test_guard_blocks_when_nodes_exist():
    c = FakeClient(get_map={"db/NODE": {"NODE": {"1": {"X": 0}}},
                            "db/ELEM": {"message": ""}})
    with pytest.raises(ModelNotEmpty):
        empty_model_guard(c, force=False, log=_quiet)


def test_guard_force_bypasses():
    c = FakeClient(get_map={"db/NODE": {"NODE": {"1": {"X": 0}}}})
    empty_model_guard(c, force=True, log=_quiet)


def test_guard_ok_when_empty():
    empty_model_guard(FakeClient(), force=False, log=_quiet)


# --- writer: matl/sect 존재 ------------------------------------------
def test_check_matl_sect_present():
    c = FakeClient(get_map={"db/MATL": {"MATL": {"1": {}}},
                            "db/SECT": {"SECT": {"1": {}}}})
    check_matl_sect(c, 1, 1, log=_quiet)


def test_check_matl_sect_missing_matl():
    c = FakeClient(get_map={"db/MATL": {"MATL": {"2": {}}},
                            "db/SECT": {"SECT": {"1": {}}}})
    with pytest.raises(MatlSectMissing) as ei:
        check_matl_sect(c, 1, 1, log=_quiet)
    assert ei.value.missing_matl == 1 and ei.value.missing_sect is None


def test_check_matl_sect_empty_collection_is_missing():
    c = FakeClient(get_map={"db/MATL": {"message": ""},
                            "db/SECT": {"message": ""}})
    with pytest.raises(MatlSectMissing):
        check_matl_sect(c, 1, 1, log=_quiet)


# --- writer: 순서/바디 --------------------------------------------
def test_write_model_puts_node_then_elem():
    m = build_model(make_params())
    c = FakeClient()
    write_model(c, m, log=_quiet)
    assert [p for p, _b in c.puts] == ["db/NODE", "db/ELEM"]
    assert c.puts[0][1] == {"Assign": m.node_assign}
    assert c.puts[1][1] == {"Assign": m.elem_assign}


def test_put_unit_first_body():
    c = FakeClient()
    put_unit(c, log=_quiet)
    assert c.puts[0][0] == "db/UNIT"


# --- CLI ----------------------------------------------------------
_FULL_GET = {
    "db/MATL": {"MATL": {"1": {}}},
    "db/SECT": {"SECT": {"1": {}}},
    "db/NODE": {"message": ""},
    "db/ELEM": {"message": ""},
}


class _NoNet:
    def __init__(self, *a, **k):
        raise AssertionError("dry-run 인데 네트워크 클라이언트 생성됨")


def _cli_fake(get_map, put_map=None):
    class _C:
        def __init__(self, mapi_key, base_url=None, timeout=30):
            self.request_log = []
            self.puts = []
            self._g = get_map
            self._p = put_map or {}

        def get(self, path):
            self.request_log.append(("GET", path, 200))
            return self._g.get(path, {"message": ""})

        def put(self, path, body):
            self.puts.append((path, body))
            self.request_log.append(("PUT", path, 200))
            return self._p.get(path, FakeResp(200, {path.rsplit("/", 1)[-1]: {}}))

        def verbs_used(self):
            return sorted({v for v, _p, _s in self.request_log})
    return _C


def _hold(monkeypatch, get_map, put_map=None):
    holder = {}
    cls = _cli_fake(get_map, put_map)

    def factory(*a, **k):
        holder["c"] = cls(*a, **k)
        return holder["c"]

    monkeypatch.setattr(cli, "WarehouseClient", factory)
    return holder


def test_make_template(tmp_path):
    out = tmp_path / "form.xlsx"
    assert cli.main(["--make-template", str(out)]) == 0
    wb = load_workbook(str(out))
    assert wb.sheetnames == [SHEET_NAME]
    assert len(list(wb.defined_names)) == 19
    assert wb[SHEET_NAME]["F5"].value is None


def test_dry_run_no_network(tmp_path, monkeypatch):
    monkeypatch.setattr(cli, "WarehouseClient", _NoNet)
    assert cli.main(["--template", make_form(tmp_path), "--dry-run"]) == 0


def test_exit_2_on_bad_span_expr(tmp_path, monkeypatch):
    monkeypatch.setattr(cli, "WarehouseClient", _NoNet)
    rc = cli.main(["--template", make_form(tmp_path, x_spans="4@"),
                   "--dry-run"])
    assert rc == 2


def test_exit_4_when_model_not_empty(tmp_path, monkeypatch):
    gm = dict(_FULL_GET)
    gm["db/NODE"] = {"NODE": {"1": {"X": 0}}}
    holder = _hold(monkeypatch, gm)
    rc = cli.main(["--template", make_form(tmp_path), "--mapi-key", "k"])
    assert rc == 4
    assert holder["c"].puts == [] and holder["c"].verbs_used() == ["GET"]


def test_exit_5_when_matl_missing(tmp_path, monkeypatch):
    gm = dict(_FULL_GET)
    gm["db/MATL"] = {"message": ""}
    holder = _hold(monkeypatch, gm)
    rc = cli.main(["--template", make_form(tmp_path), "--mapi-key", "k"])
    assert rc == 5
    assert holder["c"].puts == [] and holder["c"].verbs_used() == ["GET"]


def test_success_flow_get_and_put_only(tmp_path, monkeypatch):
    holder = _hold(monkeypatch, _FULL_GET)
    rc = cli.main(["--template", make_form(tmp_path), "--mapi-key", "k"])
    assert rc == 0
    assert holder["c"].verbs_used() == ["GET", "PUT"]
    assert [p for p, _b in holder["c"].puts] == ["db/UNIT", "db/NODE", "db/ELEM"]
