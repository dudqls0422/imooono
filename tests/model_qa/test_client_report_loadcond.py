"""클라이언트 / 리포트 / loadcond / 단위정규화 / 메시 테스트."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from tools.model_qa.checks import Finding
from tools.model_qa.client import (BaseUrlUnreachable, ReadOnlyClient,
                                   ResourceUnavailable)
from tools.model_qa.loadcond import extract_seismic, extract_wind
from tools.model_qa.report import write_report
from tools.model_qa.scan import _DIST_TO_M, _norm_nodes

from .conftest import make_ctx, node, plate


class _Resp:
    def __init__(self, status, payload):
        self.status_code = status
        self._p = payload
        self.text = str(payload)

    def json(self):
        return self._p


class _Sess:
    def __init__(self, routes):
        self.routes = routes
        self.headers = {}

    def get(self, url, timeout=None):
        for suffix, resp in self.routes.items():
            if url.endswith(suffix):
                if isinstance(resp, Exception):
                    raise resp
                return resp
        return _Resp(404, "not found")


def _client(routes):
    c = ReadOnlyClient("k")
    c._s = _Sess(routes)
    return c


def test_client_get_only_blocks_mutating():
    c = ReadOnlyClient("k")
    for m in ("post", "put", "delete"):
        with pytest.raises(RuntimeError):
            getattr(c, m)("x")


def test_client_unwrap_envelope():
    c = _client({"db/NODE": _Resp(200, {"NODE": {"1": {"X": 0}}})})
    assert c.get_collection("NODE") == {"1": {"X": 0}}


def test_client_empty_collection_is_not_error():
    c = _client({"db/POSL": _Resp(200, {"message": ""})})
    assert c.get_collection("POSL") == {}


def test_client_non200_raises_resource_unavailable():
    c = _client({"db/FOO": _Resp(500, "err")})
    with pytest.raises(ResourceUnavailable):
        c.get_collection("FOO")


def test_client_smoke_ok():
    c = _client({"db/UNIT": _Resp(200, {"UNIT": {"1": {"FORCE": "KN", "DIST": "MM"}}})})
    assert c.smoke()["DIST"] == "MM"


def test_client_smoke_fail_raises():
    import requests
    c = _client({"db/UNIT": requests.ConnectionError("down")})
    with pytest.raises(BaseUrlUnreachable):
        c.smoke()


def test_unit_mm_to_m():
    out = _norm_nodes({"1": node(1000, 2000, 3000)}, _DIST_TO_M["MM"])
    assert out["1"]["X"] == pytest.approx(1.0)
    assert out["1"]["Z"] == pytest.approx(3.0)


def test_mesh_skips_without_plates():
    from tools.model_qa.checks.mesh import m1_aspect_ratio
    ctx = make_ctx(nodes={"1": node(0, 0, 0)}, elems={})
    assert m1_aspect_ratio(ctx) == []


def test_mesh_aspect_ratio_error():
    from tools.model_qa.checks.mesh import m1_aspect_ratio
    ctx = make_ctx(
        nodes={"1": node(0, 0, 0), "2": node(20, 0, 0),
               "3": node(20, 1, 0), "4": node(0, 1, 0)},
        elems={"1": plate([1, 2, 3, 4])})
    f = m1_aspect_ratio(ctx)
    assert any(x.severity == "오류" and 1 in x.target_ids for x in f)


def test_loadcond_seismic_sseis_r_from_response_mod_factor():
    ctx = make_ctx(sseis={"5": {
        "SEIS_CODE": "KDS(41-17-00:2019)", "DESC": "s", "SCALE_FACTOR_X": 1,
        "SCALE_FACTOR_Y": 0, "ACCIDENT_TORSION": True,
        "PARAMETERS": {"SEIS_ZONE": 0, "EPA": 0.22, "SITE_CLASS": 4,
                       "IMPORTANCE_FACTOR": 1, "PERIOD_METHOD": 1,
                       "PERIOD_APPR_X": 0.58, "PERIOD_APPR_Y": 0.58,
                       "RESPONSE_MOD_FACTOR_X": 3, "RESPONSE_MOD_FACTOR_Y": 3,
                       "SDS": 0.47, "SD1": 0.38}}})
    rows = extract_seismic(ctx)
    assert rows and "X=3" in rows[0]["반응수정계수 R"]
    assert rows[0]["감쇠비"] == "미지원"
    assert rows[0]["값 출처"] == "API조회"
    assert rows[0]["지반종류"] == "S5"


def test_loadcond_seismic_splc_damping_from_spfc():
    ctx = make_ctx(
        splc={"1": {"NAME": "REX", "DIR": "XY", "COMTYPE": "CQC",
                    "aFUNCNAME": ["KDS"], "bACCECC": False}},
        spfc={"1": {"NAME": "KDS", "DRATIO": 0.05,
                    "VAL": {"IE": 1, "R_": 3, "ZONEFACTOR": 0.22}, "OPT": {"SC_": 4}}})
    rows = extract_seismic(ctx)
    assert rows and rows[0]["감쇠비"].startswith("0.05")
    assert "API조회" in rows[0]["감쇠비"]


def test_loadcond_wind_direct_parse():
    ctx = make_ctx(swind={"3": {
        "WIND_CODE": "KDS(41-12:2022)", "DESC": "",
        "PARAMETERS": {"WIND_SPEED": 28, "EXP_CATEGORY": 1, "IMPORTANCE_FACTOR": 1,
                       "GUST_FACTOR_X": 2.2, "GUST_FACTOR_Y": 2.2, "ROOF_HEIGHT": 13.62,
                       "TOPOGRAPHIC_EFFECT": {"OPT_USE": False},
                       "FORCE_COEF": {"OPT_USE": False}},
        "PROFILE": {"X_DIR": [{"PRESSURE": "0.83"}], "Y_DIR": [{"PRESSURE": "0.7"}]}}})
    rows = extract_wind(ctx)
    assert rows and rows[0]["기본풍속 Vo"] == 28
    assert rows[0]["지표면조도구분"] == "B"
    assert rows[0]["값 출처"] == "API조회"
    assert "X" in rows[0]["적용방향"] and "Y" in rows[0]["적용방향"]


def test_loadcond_wind_fallback_when_swind_empty():
    ctx = make_ctx(swind={},
                   stld={"3": {"NO": 3, "NAME": "WX", "TYPE": "W"}},
                   stor={"1": {"WIND_FLOOR_WIDTH_X": 51.6, "WIND_ECCENT_X": 7.7}})
    rows = extract_wind(ctx)
    assert rows and rows[0]["케이스명"] == "WX"
    assert "간접" in rows[0]["값 출처"] and "직접 조회 실패" in rows[0]["값 출처"]


def _mk_findings():
    fs = [Finding("G1", "중복 절점", "오류", "NODE", [1, 2], "d", "r"),
          Finding("B2", "전역 강체거동 미구속", "경고", "MODEL", [], "d", "r"),
          Finding("G7", "미접합 교차부재", "정보", "ELEM", [3], "d", "r"),
          Finding("G5", "완전 고립 절점", "경고", "NODE", list(range(1, 80)), "d", "r")]
    fs[0].whitelisted = True
    return fs


def test_report_three_sheets_and_summary_sum(tmp_path):
    ctx = make_ctx(nodes={"1": node(0, 0, 0)}, elems={},
                   stld={"1": {"NAME": "DL", "TYPE": "D"}})
    out = tmp_path / "r.xlsx"
    write_report(str(out), _mk_findings(), ctx, [], [],
                 {"n_checks": 25, "base_url": "x"})
    wb = load_workbook(out)
    assert wb.sheetnames[:3] == ["요약", "모델점검", "하중조건"]
    assert "오버플로" in wb.sheetnames

    chk = wb["모델점검"]
    nonwl_rows = sum(1 for r in chk.iter_rows(min_row=2)
                     if r[0].value and r[7].value != "Y")
    smry = {r[0].value: r[1].value for r in wb["요약"].iter_rows()}
    total = (smry.get("심각도 — 오류", 0) + smry.get("심각도 — 경고", 0)
             + smry.get("심각도 — 정보", 0))
    assert total == nonwl_rows == 3


def test_report_loadcond_has_value_source_column(tmp_path):
    ctx = make_ctx(nodes={}, elems={})
    wind = [{"케이스ID": "3", "케이스명": "WX", "값 출처": "API조회"}]
    out = tmp_path / "r2.xlsx"
    write_report(str(out), [], ctx, wind, [], {"n_checks": 25})
    wb = load_workbook(out)
    texts = [c.value for row in wb["하중조건"].iter_rows() for c in row if c.value]
    assert "값 출처" in texts and "API조회" in texts
