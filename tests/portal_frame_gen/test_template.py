"""입력 폼 생성 + 정의된 이름/폴백 파서 + 검증 규칙."""
from __future__ import annotations

import warnings

import pytest
from openpyxl import load_workbook

from tools.portal_frame_gen.geometry import build_model
from tools.portal_frame_gen.template import (SHEET_NAME, TemplateError,
                                             read_template, validate,
                                             write_template)

from .conftest import make_form, make_params


# --- validate() 규칙 (직접 호출, 폼과 무관) ------------------------------
def test_valid_params_pass():
    assert validate(make_params()) is not None


@pytest.mark.parametrize("over", [
    {"span_m": 0}, {"span_m": -1}, {"eave_height_m": 0},
    {"bay_count": 0}, {"bay_spacing_m": 0}, {"bay_spacing_m": -3},
    {"pitch_rise": 0}, {"base_fixity": "roller"}, {"frame_mode": "1D"},
])
def test_reject_rules(over):
    with pytest.raises(TemplateError):
        validate(make_params(**over))


def test_warn_steep_roof_30_to_45():
    p = make_params(pitch_rise=1.0, pitch_run=1.5)  # ~33.7°
    validate(p)
    assert p.warnings and "경사각" in p.warnings[0]


def test_reject_roof_over_45():
    with pytest.raises(TemplateError):
        validate(make_params(pitch_rise=1.0, pitch_run=0.8))  # ~51°


# --- 1. 왕복: 예시값 파싱 ------------------------------------------------
def test_roundtrip_example_values(tmp_path):
    path = make_form(tmp_path)
    r = read_template(path)
    assert r.span_m == 28.0 and isinstance(r.span_m, float)
    assert r.bay_count == 5 and isinstance(r.bay_count, int)
    assert r.pitch_rise == pytest.approx(0.15) and r.pitch_run == 1.0
    assert r.eave_height_m == 6.0 and r.bay_spacing_m == 6.0
    assert r.frame_mode == "3D" and r.base_fixity == "pinned"
    assert r.material_name == "SS275"
    assert r.column_section == "H-400x200x8x13"
    assert r.rafter_section == "H-350x175x7x11"
    assert r.eave_strut_section == "H-200x100x5.5x8"
    assert r.base_level_m == 0.0 and r.roof_angle_deg is None


# --- 2. 정의된 이름 삭제 → 고정셀 폴백 + warning ----------------------
def test_defined_name_missing_falls_back_with_warning(tmp_path):
    path = make_form(tmp_path, drop_names=["span"])
    with pytest.warns(UserWarning, match="span"):
        r = read_template(path)
    assert r.span_m == 28.0


# --- 3. 값 없음 + 이름 없음 → ValueError(필드명·셀) ------------------
def test_missing_value_and_name_raises(tmp_path):
    path = make_form(tmp_path, drop_names=["span"], span=None)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with pytest.raises(TemplateError) as ei:
            read_template(path)
    assert "span" in str(ei.value) and "F5" in str(ei.value)


# --- 4/5. sec_eave_strut 는 3D 에서만 필수 ---------------------------
def test_eave_strut_optional_in_2d(tmp_path):
    path = make_form(tmp_path, frame_mode="2D", sec_eave_strut=None)
    r = read_template(path)
    assert r.frame_mode == "2D" and r.eave_strut_section is None


def test_eave_strut_required_in_3d(tmp_path):
    path = make_form(tmp_path, frame_mode="3D", sec_eave_strut=None)
    with pytest.raises(TemplateError):
        read_template(path)


# --- 6. roof_pitch 비율/십진 ----------------------------------------
@pytest.mark.parametrize("raw,expect", [
    ("3:12", 0.25), ("3/12", 0.25), (0.15, 0.15), ("0.2", 0.2),
])
def test_roof_pitch_forms(tmp_path, raw, expect):
    path = make_form(tmp_path, roof_pitch=raw)
    r = read_template(path)
    assert r.pitch_rise == pytest.approx(expect) and r.pitch_run == 1.0


# --- 7. 잘못된 enum → 필드명 포함 ValueError ------------------------
def test_bad_frame_mode_enum(tmp_path):
    path = make_form(tmp_path, frame_mode="2.5D")
    with pytest.raises(TemplateError) as ei:
        read_template(path)
    assert "frame_mode" in str(ei.value)


# --- 8. --make-template 재실행: 정의된 이름 중복 없음 --------------
def test_regenerate_no_duplicate_defined_names(tmp_path):
    path = tmp_path / "t.xlsx"
    write_template(str(path))
    write_template(str(path))  # 같은 파일 재생성
    wb = load_workbook(str(path))
    names = list(wb.defined_names)
    assert len(names) == len(set(names)) == 11
    # 재파싱 가능(입력 채운 뒤)
    path2 = make_form(tmp_path, name="t2.xlsx")
    assert read_template(path2).span_m == 28.0


# --- 9. 회귀: dry-run 카운트가 개편 전 베이스라인과 동일 ----------
def test_regression_dry_run_counts(tmp_path):
    r = read_template(make_form(tmp_path))
    c = build_model(r).counts()
    assert c["nodes"] == 30 and c["elements"] == 34
    assert c["columns"] == 12 and c["rafters"] == 12 and c["eave_struts"] == 10
    assert c["supports"] == 12 and c["groups"] == 3


# --- 10. 데이터 유효성 존재 --------------------------------------
def test_form_has_dropdowns(tmp_path):
    path = tmp_path / "t.xlsx"
    write_template(str(path))
    wb = load_workbook(str(path))
    ws = wb[SHEET_NAME]
    formulas = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert '"2D,3D"' in formulas and '"pinned,fixed"' in formulas
