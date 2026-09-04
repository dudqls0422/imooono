"""
구조해석 결과를 캡처해 가로 A4 Excel(``해석결과 확인_EX.xlsx``)로 정리하는 코드.

원본 ``해석결과 확인.py`` 에서 다음을 확장한 수정본이다.
- GUI 에 "가새 인장력(브레이스) 캡처 포함" 체크박스 추가(기본 체크). 해제 시 TRUSS_FORCE 스텝만 건너뛴다.
- 캡처 JPG 를 ``EXPORT_DIR`` 밖 임시폴더에 만든 뒤 Excel 에 임베드하고, 종료 시 임시폴더를 삭제한다.
  따라서 ``EXPORT_DIR`` 에는 최종 ``.xlsx`` 만 남고 기존 ``model_*.jpg`` 는 건드리지 않는다.
- 캡처 파라미터(ANGLE / DISPLAY / RESULT_GRAPHIC / LOAD 등)와 뷰 로직은 원본과 100% 동일하다.

인쇄 주의: 결과 Excel 은 **이미지당 워크시트 1개**로 구성된다. 전부 인쇄하려면 인쇄
대화상자에서 '전체 통합 문서(Entire Workbook)' 를 선택해야 한다. '활성 시트만' 인쇄하면
1장만 나온다. (캡션/설명 텍스트는 넣지 않으며, 실패한 스텝만 A1 에 "[캡처 실패] ..." 표시.)

참고: ``/view/CAPTURE`` 가 이미지 바이트/base64 를 직접 반환하는지는 공식 매뉴얼 미확인이라,
이미지를 메모리로 받는 최적화는 적용하지 않았다. 원본과 동일하게 ``EXPORT_PATH`` 로 파일을
쓰게 하되 그 경로만 임시폴더로 돌리는 폴백 방식을 쓴다. 매뉴얼 확인이 되면 열려 있는 개선점이다.
"""
import itertools
import math
import os
import shutil
import tempfile
import time
import tkinter as tk
from tkinter import messagebox

import requests
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image as PILImage

BaseURL = "https://moa-engineers.midasit.com:443/gen"

# 캡처/출력 저장 경로. 스크립트 위치와 무관하게 항상 이 폴더로 최종 xlsx 를 쓴다(하드코딩 유지).
EXPORT_DIR = r"C:\Users\KYB\Desktop\해석결과 캡쳐"
try:
    os.makedirs(EXPORT_DIR, exist_ok=True)
except OSError:
    # 다른 PC / CI 등 경로가 없는 환경에서 import 자체가 깨지지 않도록 방어.
    pass

# 최종 Excel 파일명(재실행 시 덮어쓰기).
RESULT_XLSX_NAME = "해석결과 확인_EX.xlsx"

# ============ GUI로 입력받을 전역 변수 ============
MAPI_KEY = None
MATERIAL = "STL"      # 기본값: 철골
INCLUDE_TRUSS_FORCE = True   # 체크박스 기본값: 가새 인장력 캡처 포함
headers = None         # GUI 입력 후 채워짐


def start_gui():
    """MAPI-Key, 부재 종류(철골/RC), 가새 인장력 캡처 포함 여부를 입력받는 팝업창"""
    global MAPI_KEY, MATERIAL, INCLUDE_TRUSS_FORCE, headers

    root = tk.Tk()
    root.title("MIDAS 캡처 설정")
    root.geometry("380x300")
    root.resizable(False, False)

    tk.Label(root, text="MAPI-Key 입력", font=("맑은 고딕", 11, "bold")).pack(pady=(20, 5))
    key_entry = tk.Entry(root, width=48, show="")
    key_entry.pack(pady=5)

    tk.Label(root, text="부재 종류 선택", font=("맑은 고딕", 11, "bold")).pack(pady=(20, 5))
    material_var = tk.StringVar(value="STL")

    frame = tk.Frame(root)
    frame.pack(pady=5)
    tk.Radiobutton(frame, text="철골 (STL)", variable=material_var, value="STL",
                   font=("맑은 고딕", 10)).pack(side="left", padx=15)
    tk.Radiobutton(frame, text="RC", variable=material_var, value="RC",
                   font=("맑은 고딕", 10)).pack(side="left", padx=15)

    # 재질 라디오 프레임과 "확인 및 실행" 버튼 사이: 가새 인장력 캡처 포함 체크박스(기본 체크).
    truss_var = tk.BooleanVar(value=True)
    tk.Checkbutton(root, text="가새 인장력(브레이스) 캡처 포함", variable=truss_var,
                   font=("맑은 고딕", 10)).pack(pady=(15, 0))

    def on_submit():
        global MAPI_KEY, MATERIAL, INCLUDE_TRUSS_FORCE, headers
        key = key_entry.get().strip()
        if not key:
            messagebox.showwarning("입력 오류", "MAPI-Key를 입력해주세요.")
            return
        MAPI_KEY = key
        MATERIAL = material_var.get()
        INCLUDE_TRUSS_FORCE = truss_var.get()
        headers = {
            "Content-Type": "application/json",
            "MAPI-Key": MAPI_KEY
        }
        root.destroy()

    tk.Button(root, text="확인 및 실행", command=on_submit, width=22,
              bg="#4CAF50", fg="white", font=("맑은 고딕", 10, "bold")).pack(pady=25)

    root.mainloop()


# ============ 캡처 함수들 ============

def MAIN_MODLE(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "pre",
            "SET_HIDDEN": True,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "구조해석모델",
                    "LABEL_ORIENTATION": 15
                }
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    # 참고(1회 관찰용): 실 MIDAS 환경에서 /view/CAPTURE 응답 형태를 보려면 아래 주석 해제.
    #   r.json() 이 이미지 바이트/base64 를 직접 준다면 EXPORT_PATH 없이 처리 가능.
    # print("  CAPTURE(full):", r.text)
    return (r.status_code == 200, '구조해석모델')


def LOAD_DL_LL(export_path):
    body_UNIT = {
        "Assign": {"1": {"FORCE": "KN", "DIST": "M", "HEAT": "KCAL", "TEMPER": "C"}}
    }
    r1 = requests.put(BaseURL + "/db/UNIT", headers=headers, json=body_UNIT)
    print("  CAPTURE:", r1.status_code, r1.text[:200] if r1.text else "")

    body_CAP = {
        "Argument": {
            "SET_MODE": "pre",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "연직방향 하중입력",
                    "LABEL_ORIENTATION": 15
                },
                "LOAD": {
                    "CASE_SELECTION": {"TYPE": "st", "NAME": "DL"},
                    "LOAD_VALUE": {"FORMAT": "Fixed", "PLACE": 1},
                    "FLOOR_LOAD_NAME": True,
                    "NODLA_LOAD": True,
                }
            },
        }
    }
    r2 = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r2.status_code, r2.text[:200] if r2.text else "")
    return (r2.status_code == 200, '연직방향 하중입력')


def LOAD_WX(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "pre",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "X방향 풍하중입력",
                    "LABEL_ORIENTATION": 15
                },
                "LOAD": {
                    "CASE_SELECTION": {"TYPE": "st", "NAME": "WX"},
                    "LOAD_VALUE": {"FORMAT": "Fixed", "PLACE": 1},
                    "FLOOR_LOAD_NAME": True,
                    "NODLA_LOAD": True,
                }
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, 'X방향 풍하중입력')


def LOAD_WY(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "pre",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "Y방향 풍하중입력",
                    "LABEL_ORIENTATION": 15
                },
                "LOAD": {
                    "CASE_SELECTION": {"TYPE": "st", "NAME": "WY"},
                    "LOAD_VALUE": {"FORMAT": "Fixed", "PLACE": 1},
                    "FLOOR_LOAD_NAME": True,
                    "NODLA_LOAD": True,
                }
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, 'Y방향 풍하중입력')


def LOAD_EX(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "pre",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "X방향 지진하중입력",
                    "LABEL_ORIENTATION": 15
                },
                "LOAD": {
                    "CASE_SELECTION": {"TYPE": "st", "NAME": "EX"},
                    "LOAD_VALUE": {"FORMAT": "Fixed", "PLACE": 1},
                    "SEISMIC_LOAD": True,
                }
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, 'X방향 지진하중입력')


def LOAD_EY(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "pre",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "Y방향 지진하중입력",
                    "LABEL_ORIENTATION": 15
                },
                "LOAD": {
                    "CASE_SELECTION": {"TYPE": "st", "NAME": "EY"},
                    "LOAD_VALUE": {"FORMAT": "Fixed", "PLACE": 1},
                    "SEISMIC_LOAD": True,
                }
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, 'Y방향 지진하중입력')


def DEFORM_DL_LL(export_path):
    body_UNIT = {
        "Assign": {"1": {"FORCE": "KN", "DIST": "MM", "HEAT": "KCAL", "TEMPER": "C"}}
    }
    r1 = requests.put(BaseURL + "/db/UNIT", headers=headers, json=body_UNIT)
    print("  CAPTURE:", r1.status_code, r1.text[:200] if r1.text else "")

    body_COMB = {
        "Assign": {
            "9999": {
                "NO": 9999,
                "NAME": "SER",
                "ACTIVE": "ACTIVE",
                "bCB": False,
                "iTYPE": 0,
                "DESC": "",
                "vCOMB": [
                    {"ANAL": "ST", "LCNAME": "DL", "FACTOR": 1},
                    {"ANAL": "ST", "LCNAME": "LL", "FACTOR": 1}
                ]  
            }
        }
    }
    r2 = requests.put(BaseURL + "/db/LCOM-GEN", headers=headers, json=body_COMB)
    print("  CAPTURE:", r2.status_code, r2.text[:200] if r2.text else "")

    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "연직하중에 의한 변위",
                    "LABEL_ORIENTATION": 15
                },
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "Displacement Contour",
                "LOAD_CASE_COMB": {
                    "TYPE": "CB", "MINMAX": "ALL", "NAME": "SER",
                    "STEP_INDEX": 1, "TH_OPTION": "Displacement"
                },
                "COMPONENTS": {"COMP": "DXYZ", "OPT_LOCAL_CHECK": False},
                "TYPE_OF_DISPLAY": {
                    "CONTOUR": {"OPT_CHECK": True},
                    "DEFORM": {"OPT_CHECK": True, "SCALE_FACTOR": 1, "REAL_DEFORM": False},
                    "VALUES": {"OPT_CHECK": True},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"}
                },
            }
        },
    }
    r3 = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r3.status_code, r3.text[:200] if r3.text else "")
    return (r3.status_code == 200, '연직하중에 의한 변위')


def DEFORM_WX(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "풍하중에 의한 X방향 변위",
                    "LABEL_ORIENTATION": 15
                },
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "Displacement Contour",
                "LOAD_CASE_COMB": {
                    "TYPE": "ST", "MINMAX": "ALL", "NAME": "WX",
                    "STEP_INDEX": 1, "TH_OPTION": "Displacement"
                },
                "COMPONENTS": {"COMP": "DXYZ", "OPT_LOCAL_CHECK": False},
                "TYPE_OF_DISPLAY": {
                    "CONTOUR": {"OPT_CHECK": True},
                    "DEFORM": {"OPT_CHECK": True, "SCALE_FACTOR": 1, "REAL_DEFORM": False},
                    "VALUES": {"OPT_CHECK": True},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"}
                },
            }
        },
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, '풍하중에 의한 X방향 변위')


def DEFORM_WY(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "풍하중에 의한 Y방향 변위",
                    "LABEL_ORIENTATION": 15
                },
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "Displacement Contour",
                "LOAD_CASE_COMB": {
                    "TYPE": "ST", "MINMAX": "ALL", "NAME": "WY",
                    "STEP_INDEX": 1, "TH_OPTION": "Displacement"
                },
                "COMPONENTS": {"COMP": "DXYZ", "OPT_LOCAL_CHECK": False},
                "TYPE_OF_DISPLAY": {
                    "CONTOUR": {"OPT_CHECK": True},
                    "DEFORM": {"OPT_CHECK": True, "SCALE_FACTOR": 1, "REAL_DEFORM": False},
                    "VALUES": {"OPT_CHECK": True},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"}
                },
            }
        },
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, '풍하중에 의한 Y방향 변위')


def BEAM_DIAGRAMS_M(export_path):
    body_UNIT = {
        "Assign": {"1": {"FORCE": "KN", "DIST": "M", "HEAT": "KCAL", "TEMPER": "C"}}
    }
    r1 = requests.put(BaseURL + "/db/UNIT", headers=headers, json=body_UNIT)
    print("  CAPTURE:", r1.status_code, r1.text[:200] if r1.text else "")

    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": True,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "휨모멘트(ENV_STR)",
                    "LABEL_ORIENTATION": 15
                }
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "BeamDiagrams",
                "LOAD_CASE_COMB": {"TYPE": "CB", "MINMAX": "All", "NAME": f"{MATERIAL} ENV_STR", "OPT_MAXMIN_DIAGRAM": False},
                "COMPONENTS": {"PART": "Total", "COMP": "My"},
                "DISPLAY_OPTIONS": {"FIDELITY": "5Points", "FILL": "Solid", "SCALE": 1},
                "TYPE_OF_DISPLAY": {
                    "CONTOUR": {"OPT_CHECK": True},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"},
                },
            },
        }
    }
    r2 = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r2.status_code, r2.text[:200] if r2.text else "")
    return (r2.status_code == 200, '휨모멘트(ENV_STR)')


def BEAM_DIAGRAMS_S(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": True,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "전단력도(ENV_STR)",
                    "LABEL_ORIENTATION": 15
                }
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "BeamDiagrams",
                "LOAD_CASE_COMB": {"TYPE": "CB", "MINMAX": "All", "NAME": f"{MATERIAL} ENV_STR", "OPT_MAXMIN_DIAGRAM": False},
                "COMPONENTS": {"PART": "Total", "COMP": "Fz"},
                "DISPLAY_OPTIONS": {"FIDELITY": "5Points", "FILL": "Solid", "SCALE": 1},
                "TYPE_OF_DISPLAY": {
                    "CONTOUR": {"OPT_CHECK": True},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"},
                },
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, '전단력도(ENV_STR)')


def BEAM_DIAGRAMS_F(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": True,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "축력도(ENV_STR)",
                    "LABEL_ORIENTATION": 15
                }
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "BeamDiagrams",
                "LOAD_CASE_COMB": {"TYPE": "CB", "MINMAX": "All", "NAME": f"{MATERIAL} ENV_STR", "OPT_MAXMIN_DIAGRAM": False},
                "COMPONENTS": {"PART": "Total", "COMP": "Fx"},
                "DISPLAY_OPTIONS": {"FIDELITY": "5Points", "FILL": "Solid", "SCALE": 1},
                "TYPE_OF_DISPLAY": {
                    "CONTOUR": {"OPT_CHECK": True},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"},
                },
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, '축력도(ENV_STR)')


def TRUSS_FORCE(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "가새 인장력(ENV_STR)",
                    "LABEL_ORIENTATION": 15
                }
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "TrussForces",
                "LOAD_CASE_COMB": {"TYPE": "CB", "MINMAX": "All", "NAME": f"{MATERIAL} ENV_STR", "STEP_INDEX": 1},
                "COMPONENTS": {"COMP": "Tens."},
                "TYPE_OF_DISPLAY": {
                    "CONTOUR": {"OPT_CHECK": True},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"},
                    "VALUES": {"OPT_CHECK": True,  "DECIMAL_PT": 0}
                },
                "OUTPUT_SECT_LOCATION": {"OPT_I_J_MAX_ALL": "Max"},
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, '가새 인장력(ENV_STR)')




def REACTION_FORCES_STR(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "지점반력(ENV_STR)",
                    "LABEL_ORIENTATION": 15
                }
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "ReactionForces/Moments",
                "LOAD_CASE_COMB": {"TYPE": "CB", "MINMAX": "All", "NAME": f"{MATERIAL} ENV_STR"},
                "COMPONENTS": {"COMP": "Fxyz"},
                "DISPLAY_OPTIONS": {"FIDELITY": "5Points", "FILL": "Solid", "SCALE": 1},
                "TYPE_OF_DISPLAY": {
                    "VALUES": {"OPT_CHECK": True, "ARROW_SCALE_FACTOR": 0.7},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"},
                },
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, '지점반력(ENV_STR)')


def REACTION_FORCES_SER(export_path):
    body_CAP = {
        "Argument": {
            "SET_MODE": "post",
            "SET_HIDDEN": False,
            "EXPORT_PATH": export_path,
            "HEIGHT": 768,
            "WIDTH": 1280,
            "ACTIVE": {"ACTIVE_MODE": "All"},
            "ANGLE": {"HORIZONTAL": 30, "VERTICAL": 25},
            "DISPLAY": {
                "NODE": {"NODE": False, "NODE_NUMBER": False},
                "VIEW": {
                    "UCS_AXIS": False,
                    "VIEWPPORT_GIZMO": False,
                    "VIEW_POINT": False,
                    "DESCRIPTION": "지점반력(ENV_SER)",
                    "LABEL_ORIENTATION": 15
                }
            },
            "RESULT_GRAPHIC": {
                "CURRENT_MODE": "ReactionForces/Moments",
                "LOAD_CASE_COMB": {"TYPE": "CB", "MINMAX": "All", "NAME": f"{MATERIAL} ENV_SER"},
                "COMPONENTS": {"COMP": "Fxyz"},
                "DISPLAY_OPTIONS": {"FIDELITY": "5Points", "FILL": "Solid", "SCALE": 1},
                "TYPE_OF_DISPLAY": {
                    "VALUES": {"OPT_CHECK": True, "ARROW_SCALE_FACTOR": 0.7},
                    "LEGEND": {"OPT_CHECK": True, "POSITION": "right"},
                },
            },
        }
    }
    r = requests.post(BaseURL + "/view/CAPTURE", headers=headers, json=body_CAP)
    print("  CAPTURE:", r.status_code, r.text[:200] if r.text else "")
    return (r.status_code == 200, '지점반력(ENV_SER)')


# ============ GUI/OS 호출 방어 래퍼 (헤드리스 환경에서도 죽지 않게) ============

def _safe_showinfo(title, message):
    try:
        messagebox.showinfo(title, message)
    except Exception:
        pass


def _safe_showwarning(title, message):
    try:
        messagebox.showwarning(title, message)
    except Exception:
        pass


def _safe_startfile(path):
    try:
        os.startfile(path)  # Windows 전용
    except Exception:
        pass


# ============ 캡처 반복 실행 ============

_counter = itertools.count(1)

# 스텝이 캡션을 반환하기 전에 예외로 죽었을 때 쓰는 폴백 캡션.
# 정상 경로는 각 스텝이 반환한 캡션(= body 의 DISPLAY.VIEW.DESCRIPTION 원문)을 그대로 쓰고,
# 이 dict 는 예외 경로에서만 참조한다. 어느 캡션 셀에도 함수명이 노출되지 않게 하기 위함.
STEP_CAPTIONS = {
    "MAIN_MODLE": "구조해석모델",
    "LOAD_DL_LL": "연직방향 하중입력",
    "LOAD_WX": "X방향 풍하중입력",
    "LOAD_WY": "Y방향 풍하중입력",
    "LOAD_EX": "X방향 지진하중입력",
    "LOAD_EY": "Y방향 지진하중입력",
    "DEFORM_DL_LL": "연직하중에 의한 변위",
    "DEFORM_WX": "풍하중에 의한 X방향 변위",
    "DEFORM_WY": "풍하중에 의한 Y방향 변위",
    "BEAM_DIAGRAMS_M": "휨모멘트(ENV_STR)",
    "BEAM_DIAGRAMS_S": "전단력도(ENV_STR)",
    "BEAM_DIAGRAMS_F": "축력도(ENV_STR)",
    "TRUSS_FORCE": "가새 인장력(ENV_STR)",
    "REACTION_FORCES_STR": "지점반력(ENV_STR)",
    "REACTION_FORCES_SER": "지점반력(ENV_SER)",
}


def build_steps(include_truss_force):
    """실행할 캡처 스텝 함수 리스트를 반환한다. 순수 함수.

    include_truss_force=False 이면 TRUSS_FORCE(가새 인장력) 스텝만 제외한다(호출 자체를 안 함).
    나머지 순서·구성은 그대로 유지한다.
    """
    steps = [
        MAIN_MODLE, LOAD_DL_LL, LOAD_WX, LOAD_WY, LOAD_EX, LOAD_EY,
        DEFORM_DL_LL, DEFORM_WX, DEFORM_WY,
        BEAM_DIAGRAMS_M, BEAM_DIAGRAMS_S, BEAM_DIAGRAMS_F,
        TRUSS_FORCE,
        REACTION_FORCES_STR, REACTION_FORCES_SER,
    ]
    if not include_truss_force:
        steps = [s for s in steps if s is not TRUSS_FORCE]
    return steps


def next_export_path(base_dir):
    """base_dir 안에 model_001.jpg, model_002.jpg ... 순번 경로를 만든다.

    원본과 달리 저장 폴더를 인자로 받는다 — 캡처는 EXPORT_DIR 밖 임시폴더에 쓰기 위함.
    """
    idx = next(_counter)
    return os.path.join(base_dir, f"model_{idx:03d}.jpg")


def wait_for_file(path, timeout=5.0, interval=0.5):
    """path 에 크기 > 0 인 파일이 생길 때까지 최대 timeout 초(0.5s x 10) 폴링. 생기면 True."""
    deadline = time.time() + timeout
    while True:
        try:
            if os.path.getsize(path) > 0:
                return True
        except OSError:
            pass
        if time.time() >= deadline:
            return False
        time.sleep(interval)


def run_captures(temp_dir, include_truss_force):
    """스텝을 순서대로 실행해 (이미지경로 | None, 캡션) 리스트를 만든다. 네트워크 IO 담당.

    캡처 응답이 200 이고 temp_dir 에 파일이 실제로 생기면 (경로, 캡션),
    그 외(비200 / 파일 미생성 / 예외)는 (None, 캡션) 을 넣어 페이지가 누락되지 않게 한다.
    예외로 스텝이 캡션을 반환하지 못한 경우엔 STEP_CAPTIONS 폴백 캡션을 쓴다(함수명 노출 금지).
    """
    items = []
    for step in build_steps(include_truss_force):
        export_path = next_export_path(temp_dir)
        try:
            ok, caption = step(export_path)
        except Exception as exc:  # 네트워크/응답 파싱 실패, 반환값 형식 오류 등
            print(f"  [예외] {step.__name__}: {exc}")
            items.append((None, STEP_CAPTIONS.get(step.__name__, "알 수 없는 스텝")))
            continue
        if ok and wait_for_file(export_path):
            items.append((export_path, caption))
        else:
            items.append((None, caption))
    return items


def build_result_excel(items, out_path, material, orientation="landscape"):
    """캡처 결과를 **이미지당 워크시트 1개**인 Excel 로 만든다. 순수 함수(네트워크 없음).

    이 Excel 은 이미지당 시트 1개로 구성된다. 전부 인쇄하려면 인쇄 대화상자에서
    '전체 통합 문서(Entire Workbook)' 를 선택할 것. '활성 시트만' 을 인쇄하면 1장만 나온다.
    인쇄 시 이미지/텍스트는 A4 페이지 정중앙(가로·세로)에 온다
    (모든 시트 print_options.horizontalCentered / verticalCentered = True).

    items : (이미지경로 | None, 캡션) 리스트. 리스트 순서 = 시트 순서.
            성공 시트는 이미지만(셀 텍스트 없음), 실패 시트(경로 None 또는 파일 없음)는
            A1 에 빨간 '[캡처 실패] <캡션>'. 페이지 누락 없음.
    material : 레이아웃에 쓰이지 않는다(시그니처/호출부 안정성을 위해 유지).
    items 가 비면 ValueError.
    """
    if not items:
        raise ValueError("items 가 비어 있어 Excel 을 생성할 수 없습니다.")

    _ = material  # 레이아웃 미사용 (의도적으로 유지)
    TARGET_WIDTH_PX = 1180
    fail_font = Font(bold=True, color="FFFF0000")

    wb = Workbook()

    for i, (img_path, caption) in enumerate(items):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = f"{i + 1:02d}"  # "01", "02", ... 캡션/설명은 탭에 넣지 않는다

        # 모든 시트 동일: 가로 A4, 시트 1개 = 물리 페이지 1장 (fitToHeight=1 로 구조적 보장).
        ws.page_setup.paperSize = 9  # 9 = A4
        ws.page_setup.orientation = orientation
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5,
                                      header=0.2, footer=0.2)
        # 헤더/푸터/페이지번호 없음 (기본이 비어 있어도 방어적으로 비운다).
        for part in ("oddHeader", "oddFooter", "evenHeader", "evenFooter"):
            hf = getattr(ws, part)
            hf.left.text = ""
            hf.center.text = ""
            hf.right.text = ""

        # 인쇄 시 내용을 A4 페이지 가로·세로 정중앙에 (여백 재분배만 하므로 fitTo 와 충돌 없음).
        ws.print_options.horizontalCentered = True
        ws.print_options.verticalCentered = True

        if img_path and os.path.exists(img_path):
            try:
                with PILImage.open(img_path) as im:
                    iw, ih = im.size
            except Exception:
                iw, ih = 1280, 768
            if not iw:
                iw, ih = 1280, 768
            xl_img = XLImage(img_path)
            xl_img.width = TARGET_WIDTH_PX
            xl_img.height = int(round(TARGET_WIDTH_PX * ih / iw))
            ws.add_image(xl_img, "A1")
            # 구버전 Excel 에서도 중앙정렬 기준 범위가 잡히도록 이미지 풋프린트를 print_area 로 명시.
            # 기본 열폭 ~64px, 기본 행높이 ~20px 기준 + 여유 1.
            cols = math.ceil(xl_img.width / 64) + 1
            rows = math.ceil(xl_img.height / 20) + 1
            ws.print_area = f"A1:{get_column_letter(cols)}{rows}"
        else:
            cell = ws["A1"]
            cell.value = f"[캡처 실패] {caption}"
            cell.font = fail_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    out_dir = os.path.dirname(out_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(out_path)


def main():
    """임시폴더 캡처 -> Excel 생성 -> 임시폴더 삭제(finally). MAPI_KEY/headers 는 GUI 에서 채워진 상태여야 함."""
    temp_dir = tempfile.mkdtemp(prefix="해석결과캡처_")
    try:
        items = run_captures(temp_dir, INCLUDE_TRUSS_FORCE)
        success = sum(1 for p, _ in items if p is not None)
        out_path = os.path.join(EXPORT_DIR, RESULT_XLSX_NAME)

        if success == 0:
            print("성공한 캡처가 0건이라 Excel 파일을 생성하지 않습니다.")
            _safe_showwarning("캡처 실패", "성공한 캡처가 없어 Excel 파일을 만들지 않았습니다.")
            return

        build_result_excel(items, out_path, MATERIAL)
        print(f"\n완료: {out_path}  (총 {len(items)} 페이지, 캡처 성공 {success}건)")
        _safe_showinfo("완료", f"{RESULT_XLSX_NAME} 저장 완료\n총 {len(items)} 페이지 (성공 {success}건)")
        _safe_startfile(out_path)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


if __name__ == "__main__":
    start_gui()          # MAPI-Key, 재질, 가새 인장력 포함 여부 입력창
    if MAPI_KEY:
        main()
    else:
        print("MAPI-Key가 입력되지 않아 종료합니다.")